import sys, json, zipfile, re
from docx import Document
p = sys.argv[1]
d = Document(p)
z = zipfile.ZipFile(p)
xml = z.read('word/document.xml').decode('utf8')
print('sections:', len(d.sections), '| paragraphs:', len(d.paragraphs), '| tables:', len(d.tables), '| inline images:', len(d.inline_shapes))
print('w:bidi paragraphs:', xml.count('<w:bidi/>') + xml.count('<w:bidi '), '| w:rtl runs:', xml.count('<w:rtl/>') + xml.count('<w:rtl '), '| visuallyRightToLeft tables:', xml.count('w:bidiVisual'))
s = d.sections[0]
print('page1 size pt:', round(s.page_width.pt), 'x', round(s.page_height.pt), 'margins pt L/R/T/B:', [round(x.pt) for x in (s.left_margin, s.right_margin, s.top_margin, s.bottom_margin)])
quran = json.load(open('data/quran-uthmani.json', encoding='utf8'))
def verse(sn, a): return next(v[2] for v in quran if v[0]==sn and v[1]==a)
txt = '\n'.join(par.text for par in d.paragraphs)
ok255 = verse(2,255) in txt
ok112 = ' '.join(verse(112,a) for a in (1,2,3,4)) in txt
print('Ayat al-Kursi exact in docx:', ok255, '| Al-Ikhlas exact in docx:', ok112)
print('unmapped markers left:', txt.count('�'))
for par in d.paragraphs[:4]: print('  P:', par.text[:90])
if d.tables:
    t = d.tables[0]
    for r in t.rows: print('  ROW:', ' | '.join(c.text for c in r.cells))
